"""
document_loader.py
-------------------
Extracts raw text from uploaded resume files (PDF, DOCX, TXT).
Every function fails gracefully and returns an (text, error) tuple so the
rest of the application never crashes on a malformed or unsupported file.
"""

import io


def extract_text_from_pdf(file_bytes: bytes):
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                continue
        text = "\n".join(text_parts).strip()
        if not text:
            return "", "The PDF appears to contain no extractable text (it may be a scanned image)."
        return text, None
    except Exception as e:
        return "", f"Could not read PDF file: {e}"


def extract_text_from_docx(file_bytes: bytes):
    try:
        import docx

        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        text = "\n".join(paragraphs).strip()
        if not text:
            return "", "The DOCX file appears to be empty."
        return text, None
    except Exception as e:
        return "", f"Could not read DOCX file: {e}"


def extract_text_from_txt(file_bytes: bytes):
    try:
        text = file_bytes.decode("utf-8", errors="ignore").strip()
        if not text:
            return "", "The TXT file appears to be empty."
        return text, None
    except Exception as e:
        return "", f"Could not read TXT file: {e}"


def extract_text(uploaded_file):
    """
    Accepts a Streamlit UploadedFile object and dispatches to the right
    extractor based on its extension. Returns (text, error_message_or_None).
    """
    if uploaded_file is None:
        return "", "No file was uploaded."

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif filename.endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    else:
        return "", "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
